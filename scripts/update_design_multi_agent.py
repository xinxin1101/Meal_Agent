from docx import Document


path = "MealPilot_最终设计方案_v1.1.docx"
doc = Document(path)
doc.add_paragraph("17. 用户确认的三角色 Agent 编排", style="Heading 1")
doc.add_paragraph(
    "本项目采用受限的三角色编排：规划 Agent、确定性验证器和说明 Agent。该编排不替代 OR-Tools 与 Decimal 验证，而是将自然语言理解、候选检索和用户沟通与权威计算分离。"
)
doc.add_paragraph("规划 Agent", style="Heading 2")
doc.add_paragraph(
    "负责解析用户的显式需求、区分硬约束与软偏好并生成检索策略；只能调用受注册的检索与 Solver 工具。它不得推断健康目标、修改过敏原/忌口，不得直接选择最终菜谱或份量。"
)
doc.add_paragraph("确定性验证器", style="Heading 2")
doc.add_paragraph(
    "不是 LLM Agent。它以版本化菜谱和营养数据为依据，使用 Decimal 复算过敏原、禁忌、营养和时间，只输出通过/不通过及机器可读原因。其结论高于任何模型输出。"
)
doc.add_paragraph("说明 Agent", style="Heading 2")
doc.add_paragraph(
    "仅在方案通过确定性验证后运行，只接收已验证的餐次、份量和汇总数据，用于生成用户可读说明、替代建议和协商文案。它不能增加菜谱、修改约束、重算数值或改变已验证方案。"
)
doc.add_paragraph("故障与降级", style="Heading 2")
doc.add_paragraph(
    "规划 Agent 的模型不可用、超时或输出不合规时，系统回退到白名单规则解析；说明 Agent 不可用时省略自然语言说明。无论何种降级，检索硬过滤、Solver 与确定性验证都必须继续独立工作。"
)
doc.save(path)
